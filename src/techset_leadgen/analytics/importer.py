import codecs
import csv
import hashlib
import math
import os
import unicodedata
import logging
import zipfile
from logging.handlers import RotatingFileHandler
from datetime import date, datetime
from techset_leadgen.domain.lead_stages import normalize_lead_status

from techset_leadgen.database.local_db import get_connection, get_data_dir, get_logs_dir
from techset_leadgen.utils.formatting import pluralize

_openpyxl = None

def _get_openpyxl():
    global _openpyxl
    if _openpyxl is None:
        from openpyxl import load_workbook
        _openpyxl = load_workbook
    return _openpyxl

_LOGGER = None
MAX_IMPORT_FILE_SIZE_MB = 50
MAX_IMPORT_FILE_SIZE_BYTES = MAX_IMPORT_FILE_SIZE_MB * 1024 * 1024

REQUIRED_LOGICAL_COLS = {"client_name", "client_id", "product_type", "price"}

SOURCE_NORMALIZATION = {
    "facebook": "meta_ads",
    "fb": "meta_ads",
    "meta": "meta_ads",
    "instagram": "meta_ads",
    "ig": "meta_ads",
    "messenger": "meta_ads",
    "audience_network": "meta_ads",

    "google": "google_ads",
    "google_ads": "google_ads",
    "adwords": "google_ads",
    "gads": "google_ads",

    "linkedin": "linkedin_ads",
    "linkedin_ads": "linkedin_ads",

    "tiktok": "tiktok_ads",
    "tiktok_ads": "tiktok_ads",

    "csv": "manual_import",
    "manual": "manual_import",
}

REQUIRED_LEADS_COLS = {"created_at", "source"}
DEFAULT_LEAD_STATUS = "new"

REQUIRED_MARKETING_SPEND_COLS = {"source", "amount"}
OPTIONAL_MARKETING_SPEND_COLS = {"date"}

MARKETING_SPEND_HEADER_ALIASES = {
    "source": [
        "source", "platform", "channel", "fuente", "plataforma", "canal",
    ],
    "amount": [
        "amount", "spend", "investment", "importe", "gasto",
        "inversion", "inversión", "monto", "cost", "coste",
    ],
    "date": [
        "date", "fecha", "period", "periodo", "month", "mes",
    ],
}
OPTIONAL_LEADS_COLS = {"external_id", "status", "meeting_date", "client_id"}

LEADS_HEADER_ALIASES = {
    "external_id": [
        "leadgen_id",
        "gclid",
        "lead_id",
        "external_id",
        "campaign_id",
    ],
    "created_at": [
        "created_time",
        "conversion_time",
        "date",
        "created_at",
        "datetime",
    ],
    "source": [
        "source",
        "platform",
        "origin",
        "channel",
    ],
    "status": [
        "status",
        "state",
        "lead_status",
        "campaign_status",
    ],
    "meeting_date": [
        "meeting_date",
        "appointment_date",
        "scheduled_date",
    ],
    "client_id": [
        "client_id",
        "cliente_id",
        "identificador",
    ],
}

_SHARED_LEAD_SPEND_HEADERS = {
    "date", "datetime", "source", "platform", "origin", "channel",
    "campaign_id", "campaign_status",
}
_LEAD_ONLY_HEADERS = {
    alias
    for aliases in LEADS_HEADER_ALIASES.values()
    for alias in aliases
    if alias not in _SHARED_LEAD_SPEND_HEADERS
}

HEADER_ALIASES = {
    "cliente": "client_name",
    "nombre cliente": "client_name",
    "nombre": "client_name",
    "client_name": "client_name",
    "identificador del cliente": "client_id",
    "cliente id": "client_id",
    "id cliente": "client_id",
    "identificador": "client_id",
    "client_id": "client_id",
    "tipo de producto": "product_type",
    "producto": "product_type",
    "tipo producto": "product_type",
    "tipo": "product_type",
    "product_type": "product_type",
    "precio": "price",
    "importe": "price",
    "price": "price",
    "fecha": "date",
    "date": "date",
}


def _normalise_header(h: str) -> str:
    if not h:
        return ""

    h = unicodedata.normalize("NFKD", h)

    h = h.replace("\ufeff", "").replace("\xa0", " ")

    return " ".join(h.strip().lower().split())

def _cell_number_to_text(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            return value
        if value.is_integer():
            return str(int(value))
    return str(value)


def _validate_required_string(field: str, value) -> str:
    if value is None:
        raise ValueError(f"The required field '{field}' is missing")

    value = _cell_number_to_text(value)

    if not isinstance(value, str):
        raise ValueError(
            f"The required field '{field}' must be text, "
            f"got: {type(value).__name__}"
        )

    normalized = value.strip()

    if not normalized:
        raise ValueError(f"The required field '{field}' is empty")

    return normalized

def _validate_positive_price(field: str, value) -> float:
    if value is None:
        raise ValueError(f"The field '{field}' is required")

    if isinstance(value, str):
        value = value.strip()
        if not value:
            raise ValueError(f"The field '{field}' is empty")

    if isinstance(value, bool):
        raise ValueError(
            f"The field '{field}' must be a number, got: {value!r}"
        )

    try:
        price = float(value)
    except (TypeError, ValueError):
        raise ValueError(
            f"The field '{field}' must be a number, got: {value!r}"
        )

    if not math.isfinite(price):
        raise ValueError(
            f"The field '{field}' must be a finite number, got: {value!r}"
        )

    if price <= 0:
        raise ValueError(
            f"The field '{field}' must be greater than 0, got: {price}"
        )

    return price


def _build_header_map(headers):
    logical_to_real = {}
    for real in headers:
        if real is None:
            continue
        key = _normalise_header(real)
        if key in HEADER_ALIASES:
            logical = HEADER_ALIASES[key]
            if logical not in logical_to_real:
                logical_to_real[logical] = real
    return logical_to_real


def _check_file_size(file_path: str) -> None:
    size = os.path.getsize(file_path)
    if size > MAX_IMPORT_FILE_SIZE_BYTES:
        raise ValueError(
            f"The file exceeds the maximum allowed size "
            f"({MAX_IMPORT_FILE_SIZE_MB} MB)"
        )

def _get_importer_logger() -> logging.Logger:
    global _LOGGER
    if _LOGGER:
        return _LOGGER

    logger = logging.getLogger("fdt.importer")
    logger.setLevel(logging.INFO)

    log_file = get_logs_dir() / "importer.log"

    handler = RotatingFileHandler(
        str(log_file),
        maxBytes=1_000_000,
        backupCount=3,
        encoding="utf-8",
    )

    formatter = logging.Formatter(
        "%(asctime)s | %(levelname)s | %(message)s"
    )
    handler.setFormatter(formatter)

    logger.addHandler(handler)
    logger.propagate = False

    _LOGGER = logger
    return logger
def _parse_date(value: object, default_today: str | None) -> str | None:
    if value is None or str(value).strip() == "":
        return default_today

    if isinstance(value, bool):
        raise ValueError(f"Unrecognized date format: {value}")
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()

    s = str(value).strip()

    formats = [
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%Y/%m/%d",
        "%d-%m-%Y",
        "%d/%m/%y",
        "%Y-%m-%d %H:%M:%S",
        "%d/%m/%Y %H:%M:%S",
    ]

    for fmt in formats:
        try:
            dt = datetime.strptime(s, fmt)
            return dt.date().isoformat()
        except ValueError:
            pass

    if len(s) >= 10:
        try:
            return datetime.fromisoformat(s).date().isoformat()
        except ValueError:
            pass

    if isinstance(value, (int, float)) and math.isfinite(value):
        try:
            base = date(1899, 12, 30)
            return date.fromordinal(base.toordinal() + int(value)).isoformat()
        except (ValueError, OverflowError):
            pass

    raise ValueError(f"Unrecognized date format: {value}")


def _detect_csv_encoding(file_path: str) -> str:
    decoder = codecs.getincrementaldecoder("utf-8")()
    with open(file_path, "rb") as fb:
        try:
            while chunk := fb.read(1 << 20):
                decoder.decode(chunk)
            decoder.decode(b"", final=True)
        except UnicodeDecodeError:
            return "cp1252"
    return "utf-8-sig"


def _export_failed_rows(failed_rows: list, prefix: str, logger: logging.Logger) -> str | None:
    if not failed_rows:
        return None

    fieldnames: list[str] = []
    normalized_rows = []
    for row in failed_rows:
        clean = {}
        for key, value in row.items():
            name = "_extra_columns" if key is None else str(key)
            if isinstance(value, list):
                value = " | ".join("" if v is None else str(v) for v in value)
            clean[name] = value
            if name not in fieldnames:
                fieldnames.append(name)
        normalized_rows.append(clean)

    failed_path = get_data_dir() / f"{prefix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    try:
        with open(failed_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames, restval="")
            writer.writeheader()
            writer.writerows(normalized_rows)
    except (OSError, csv.Error):
        logger.exception("Could not export the rejected rows to %s", failed_path)
        return None

    logger.info("Rejected rows exported to %s", failed_path)
    return str(failed_path)


def _read_csv_rows(file_path: str):
    encoding = _detect_csv_encoding(file_path)
    f = open(file_path, "r", encoding=encoding, errors="replace", newline="")

    sample = f.read(4096)
    f.seek(0)
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel

    reader = csv.DictReader(f, dialect=dialect)
    headers = reader.fieldnames or []

    def row_generator():
        try:
            for row in reader:
                yield row
        finally:
            f.close()

    return headers, row_generator()


def _read_excel_rows(file_path: str):
    load_workbook = _get_openpyxl()
    from openpyxl.utils.exceptions import InvalidFileException
    try:
        wb = load_workbook(file_path, read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, EOFError) as exc:
        raise ValueError(
            "The Excel file is damaged or is not a valid .xlsx file"
        ) from exc
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)

    try:
        headers = next(rows_iter)
    except StopIteration:
        wb.close()
        return [], []

    headers = [str(h) if h is not None else "" for h in headers]

    def row_generator():
        try:
            for r in rows_iter:
                row_dict = {}
                for idx, val in enumerate(r):
                    if idx < len(headers):
                        row_dict[headers[idx]] = val
                yield row_dict
        finally:
            wb.close()

    return headers, row_generator()


def read_table(file_path: str, *, max_rows: int) -> tuple[list[str], list[dict]]:
    _check_file_size(file_path)
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".csv":
        headers, row_gen = _read_csv_rows(file_path)
    elif ext in (".xlsx", ".xlsm"):
        headers, row_gen = _read_excel_rows(file_path)
    else:
        raise ValueError("Choose a .csv or .xlsx file")
    rows: list[dict] = []
    try:
        for raw in row_gen:
            if len(rows) >= max_rows:
                raise ValueError(f"The file has more than {max_rows} rows. Split it into smaller files.")
            rows.append({_normalise_header(str(key)): value for key, value in raw.items() if key is not None})
    finally:
        close = getattr(row_gen, "close", None)
        if close is not None:
            close()
    return [_normalise_header(str(header)) for header in headers], rows


def _begin_transaction(conn) -> None:
    conn.execute("BEGIN")

def import_csv_to_db(file_path: str, dry_run: bool = False):
    import time
    start_time = time.perf_counter()
    logger = _get_importer_logger()
    logger.info("Sales import started: %s", file_path)

    try:
        _check_file_size(file_path)
        ext = os.path.splitext(file_path)[1].lower()

        if ext == ".csv":
            headers, raw_rows = _read_csv_rows(file_path)
        elif ext in (".xlsx", ".xlsm"):
            headers, raw_rows = _read_excel_rows(file_path)
        else:
            return False, f"Unsupported format: {ext}. Use CSV or XLSX."

        if not headers:
            return False, "The file has no column headers."

        header_map = _build_header_map(headers)

        missing = REQUIRED_LOGICAL_COLS - set(header_map.keys())
        if missing:
            missing_str = ",".join(sorted(missing))
            found_str = ",".join(sorted(header_map.keys()))
            return False, (
                f"Missing required columns: {missing_str}\n"
                f"Detected columns: {found_str}"
            )

        today_str = date.today().isoformat()
        used_default_date = False
        valid_rows = []
        errors = []
        failed_rows = []

        has_date_col = "date" in header_map

        for i, raw in enumerate(raw_rows, start=2):
            try:
                client_name = _validate_required_string(
                    "client_name", raw.get(header_map["client_name"])
                )
                client_id = _validate_required_string(
                    "client_id", raw.get(header_map["client_id"])
                )
                product_type = _validate_required_string(
                    "product_type", raw.get(header_map["product_type"])
                )
                price = _validate_positive_price(
                    "price", raw.get(header_map["price"])
                )

                if has_date_col:
                    date_raw = raw.get(header_map["date"])
                    parsed_date = _parse_date(date_raw, today_str)
                else:
                    parsed_date = today_str
                    used_default_date = True

                valid_rows.append(
                    (client_name, client_id, product_type, price, parsed_date)
                )

            except Exception as e:
                msg = f"Row {i}: {e}"
                errors.append(msg)
                logger.warning(msg)

                failed_rows.append({
                    "row": i,
                    "error": str(e),
                    **raw
                })

        if not valid_rows:
            return False, "No valid rows to import."

        conn = get_connection()
        try:
            clients_in_file = {}
            for client_name, client_id, *_ in valid_rows:
                if client_id not in clients_in_file:
                    clients_in_file[client_id] = client_name

            _begin_transaction(conn)
            cursor = conn.cursor()

            existing_ids = set()
            if clients_in_file:
                cursor.execute(
                    "SELECT id FROM clients WHERE id IN ({})".format(
                        ",".join("?" * len(clients_in_file))
                    ),
                    list(clients_in_file.keys())
                )
                existing_ids = {row[0] for row in cursor.fetchall()}

            missing_clients = [
                (cid, cname, "-", None, None, None)
                for cid, cname in clients_in_file.items()
                if cid not in existing_ids
            ]

            if missing_clients:
                cursor.executemany(
                    """
                    INSERT INTO clients (id, name, cif, address, contact_person, email)
                    VALUES (?, ?, ?, ?, ?, ?)
                    """,
                    missing_clients
                )
                logger.info("Inserted %d new clients", len(missing_clients))

            cursor.executemany(
                """
                INSERT INTO sales (client_name, client_id, product_type, price, date)
                VALUES (?, ?, ?, ?, ?)
                """,
                valid_rows,
            )
            logger.info("Inserted %d sales", len(valid_rows))

            if dry_run:
                conn.rollback()
                logger.info("DRY RUN: transaction rolled back, nothing saved")
            else:
                conn.commit()
                logger.info("Transaction committed")

            duration = time.perf_counter() - start_time

            if not dry_run:
                try:
                    cursor.execute(
                        """
                        INSERT INTO import_metadata
                            (imported_at, file_name, total_rows, valid_rows, rejected_rows, duration_seconds)
                        VALUES (?, ?, ?, ?, ?, ?)
                        """,
                        (
                            datetime.now().isoformat(timespec="seconds"),
                            os.path.basename(file_path),
                            len(valid_rows) + len(errors),
                            len(valid_rows),
                            len(errors),
                            duration,
                        )
                    )
                    conn.commit()
                except Exception:
                    logger.exception("Could not save import metadata")

            logger.info(
                "Duration=%.2fs | rows=%s | rows/s=%.1f",
                duration,
                len(valid_rows),
                len(valid_rows) / duration if duration else 0,
            )
            logger.info(
                "Import finished | valid=%s | rejected=%s | default_date=%s | dry_run=%s",
                len(valid_rows),
                len(errors),
                used_default_date,
                dry_run,)

        except Exception:
            conn.rollback()
            raise

        finally:
            conn.close()


        if dry_run:
            msg = f"OK {pluralize(len(valid_rows), 'valid record')} analyzed."
        else:
            msg = f"OK {pluralize(len(valid_rows), 'record')} imported."
        _export_failed_rows(failed_rows, "failed_rows", logger)
        if errors:
            msg += f" ⚠ {pluralize(len(errors), 'row')} rejected."
        if used_default_date:
            msg += " ⚠ No date column: today's date was used."
        if dry_run:
            msg += "\nDRY RUN: no data was saved to the database."

        return True, msg

    except Exception as e:
        logger.exception("Sales import failed")
        return False, f"Could not process the file: {e}"


def _build_leads_header_map(headers):
    logical_to_real = {}
    for real in headers:
        if real is None:
            continue
        key = _normalise_header(real)

        for logical_name, aliases_list in LEADS_HEADER_ALIASES.items():
            normalized_aliases = [_normalise_header(a) for a in aliases_list]
            if key in normalized_aliases:
                if logical_name not in logical_to_real:
                    logical_to_real[logical_name] = real
                break

    return logical_to_real


def _build_marketing_spend_header_map(headers):
    logical_to_real = {}
    for real in headers:
        if real is None:
            continue
        key = _normalise_header(real)

        for logical_name, aliases_list in MARKETING_SPEND_HEADER_ALIASES.items():
            normalized_aliases = [_normalise_header(a) for a in aliases_list]
            if key in normalized_aliases:
                if logical_name not in logical_to_real:
                    logical_to_real[logical_name] = real
                break

    return logical_to_real


def _validate_optional_string(field: str, value) -> str | None:
    if value is None or value == "":
        return None

    if not isinstance(value, str):
        s = str(value).strip() if value is not None else None
        if not s:
            return None
        return s

    normalized = value.strip()
    if not normalized:
        return None

    return normalized


def _normalize_source(source: str) -> str:
    if not source:
        return source

    source_key = source.lower().strip()

    for suffix in [" lead ads", " lead form", " ads", "_ads", "-ads"]:
        if source_key.endswith(suffix):
            source_key = source_key[:-len(suffix)].strip()

    normalized = SOURCE_NORMALIZATION.get(source_key, source_key)

    if normalized == source_key and source_key not in SOURCE_NORMALIZATION:
        logger = _get_importer_logger()
        logger.warning(
            f"Unknown marketing source detected during import: '{source}' -> '{source_key}'"
        )

    return normalized


def inspect_import_file(file_path: str) -> dict:
    _INVALID = {
        "valid": False,
        "detected_type": "invalid",
        "headers": [],
        "missing_sales": [],
        "missing_leads": [],
        "missing_marketing_spend": [],
    }

    try:
        _check_file_size(file_path)

        ext = os.path.splitext(file_path)[1].lower()
        if ext == ".csv":
            headers, row_gen = _read_csv_rows(file_path)
        elif ext in (".xlsx", ".xlsm"):
            headers, row_gen = _read_excel_rows(file_path)
        else:
            return {**_INVALID, "missing_sales": ["unsupported file format"], "missing_leads": ["unsupported file format"]}

        if not headers:
            try:
                for _ in row_gen:
                    pass
            except Exception:
                pass
            return {**_INVALID, "missing_sales": ["no column headers"], "missing_leads": ["no column headers"]}

        sales_map     = _build_header_map(headers)
        missing_sales = sorted(REQUIRED_LOGICAL_COLS - set(sales_map.keys()))
        fits_sales    = len(missing_sales) == 0

        leads_map     = _build_leads_header_map(headers)
        missing_leads = sorted(REQUIRED_LEADS_COLS - set(leads_map.keys()))
        fits_leads    = len(missing_leads) == 0

        spend_map             = _build_marketing_spend_header_map(headers)
        missing_marketing_spend = sorted(REQUIRED_MARKETING_SPEND_COLS - set(spend_map.keys()))
        fits_marketing_spend  = len(missing_marketing_spend) == 0

        has_lead_only_header = any(
            _normalise_header(header) in _LEAD_ONLY_HEADERS for header in headers if header
        )

        if fits_sales and fits_leads:
            detected_type = "mixed"
        elif fits_sales:
            detected_type = "sales"
        elif fits_marketing_spend and not has_lead_only_header:
            detected_type = "marketing_spend"
        elif fits_leads:
            detected_type = "leads"
        elif fits_marketing_spend:
            detected_type = "marketing_spend"
        else:
            detected_type = "invalid"

        row_count = 0
        try:
            for _ in row_gen:
                row_count += 1
        except Exception:
            pass

        return {
            "valid":          detected_type != "invalid",
            "detected_type":  detected_type,
            "headers":        list(headers),
            "row_count":      row_count,
            "valid_rows":     row_count,
            "invalid_rows":   0,
            "missing_sales":  missing_sales,
            "missing_leads":  missing_leads,
            "missing_marketing_spend": missing_marketing_spend,
        }

    except (ValueError, OSError, csv.Error) as exc:
        return {**_INVALID, "missing_sales": [str(exc)], "missing_leads": [str(exc)]}


def _insert_lead_row(cursor, row: dict) -> bool:
    cursor.execute(
        """
        INSERT OR IGNORE INTO leads
            (id, external_id, created_at, source, status, meeting_date, client_id)
        VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        (
            row["id"],
            row["external_id"],
            row["created_at"],
            row["source"],
            row["status"],
            row["meeting_date"],
            row["client_id"],
        ),
    )
    return cursor.rowcount > 0


def _lead_exists_for_source(cursor, external_id: str, source: str | None) -> bool:
    cursor.execute(
        "SELECT 1 FROM leads WHERE external_id = ? AND source = ? LIMIT 1",
        (external_id, source),
    )
    return cursor.fetchone() is not None


def import_leads_to_db(file_path: str, dry_run: bool = False):
    import time

    start_time = time.perf_counter()
    logger = _get_importer_logger()
    logger.info("Lead import started: %s", file_path)

    try:
        _check_file_size(file_path)
        ext = os.path.splitext(file_path)[1].lower()

        if ext == ".csv":
            headers, raw_rows = _read_csv_rows(file_path)
        elif ext in (".xlsx", ".xlsm"):
            headers, raw_rows = _read_excel_rows(file_path)
        else:
            return False, f"Unsupported format: {ext}. Use CSV or XLSX."

        if not headers:
            return False, "The file has no column headers."

        header_map = _build_leads_header_map(headers)

        missing = REQUIRED_LEADS_COLS - set(header_map.keys())
        if missing:
            missing_str = ",".join(sorted(missing))
            found_str = ",".join(sorted(header_map.keys()))
            return False, (
                f"Missing required columns: {missing_str}\n"
                f"Detected columns: {found_str}"
            )

        valid_rows = []
        errors = []
        failed_rows = []
        duplicates_ignored = 0
        unlinked_leads = 0
        key_occurrences: dict[str, int] = {}

        for i, raw in enumerate(raw_rows, start=2):
            try:
                raw_source = raw.get(header_map["source"])

                created_at_raw = raw.get(header_map["created_at"])
                created_at = _parse_date(created_at_raw, date.today().isoformat())

                source = _validate_required_string(
                    "source", raw_source
                )

                source = _normalize_source(source)

                external_id = None
                if "external_id" in header_map:
                    ext_id_raw = raw.get(header_map["external_id"])
                    external_id = _validate_optional_string("external_id", ext_id_raw)

                status = DEFAULT_LEAD_STATUS
                if "status" in header_map:
                    status_raw = raw.get(header_map["status"])
                    status = normalize_lead_status(
                        _validate_optional_string("status", status_raw), DEFAULT_LEAD_STATUS
                    )

                meeting_date = None
                if "meeting_date" in header_map:
                    meeting_raw = raw.get(header_map["meeting_date"])
                    if meeting_raw:
                        try:
                            meeting_date = _parse_date(meeting_raw, None)
                        except ValueError:
                            meeting_date = None

                client_id = None
                if "client_id" in header_map:
                    client_id_raw = raw.get(header_map["client_id"])
                    client_id = _validate_optional_string("client_id", client_id_raw)

                if external_id:
                    lead_id = external_id
                else:
                    base_key = f"{client_id or ''}_{created_at}_{source}"
                    occurrence = key_occurrences.get(base_key, 0)
                    key_occurrences[base_key] = occurrence + 1
                    _key = base_key if occurrence == 0 else f"{base_key}#{occurrence}"
                    lead_id = hashlib.sha1(_key.encode(), usedforsecurity=False).hexdigest()

                valid_rows.append({
                    "id": lead_id,
                    "external_id": external_id,
                    "created_at": created_at,
                    "source": source,
                    "status": status,
                    "meeting_date": meeting_date,
                    "client_id": client_id,
                })

            except Exception as e:
                msg = f"Row {i}: {e}"
                errors.append(msg)
                logger.warning(msg)

                failed_rows.append({
                    "row": i,
                    "error": str(e),
                    **raw
                })

        if not valid_rows:
            return False, "No valid rows to import."

        conn = get_connection()
        try:
            _begin_transaction(conn)
            cursor = conn.cursor()

            client_ids_in_file = {
                row["client_id"] 
                for row in valid_rows 
                if row["client_id"] is not None
            }

            existing_client_ids = set()
            if client_ids_in_file:
                placeholders = ",".join("?" * len(client_ids_in_file))
                cursor.execute(
                    f"SELECT id FROM clients WHERE id IN ({placeholders})",
                    list(client_ids_in_file)
                )
                existing_client_ids = {row[0] for row in cursor.fetchall()}

            invalid_client_ids = client_ids_in_file - existing_client_ids
            if invalid_client_ids:
                logger.warning(
                    "client_ids not found in the database (lead imported without a client): %s",
                    invalid_client_ids,
                )

            insert_count = 0
            for row in valid_rows:
                if row["client_id"] is not None and row["client_id"] not in existing_client_ids:
                    row["client_id"] = None
                    unlinked_leads += 1

                inserted = _insert_lead_row(cursor, row)
                if (
                    not inserted
                    and row["external_id"]
                    and not _lead_exists_for_source(cursor, row["external_id"], row["source"])
                ):
                    inserted = _insert_lead_row(
                        cursor, {**row, "id": f"{row['source']}:{row['external_id']}"}
                    )

                if inserted:
                    insert_count += 1
                else:
                    duplicates_ignored += 1

            logger.info(
                "Leads: %d new, %d duplicates skipped",
                insert_count,
                duplicates_ignored,
            )

            if dry_run:
                conn.rollback()
                logger.info("DRY RUN: transaction rolled back, nothing saved")
            else:
                conn.commit()
                logger.info("Lead transaction committed")

            duration = time.perf_counter() - start_time

            if not dry_run:
                try:
                    cursor.execute(
                        """
                        INSERT INTO import_metadata
                            (imported_at, file_name, total_rows, valid_rows, rejected_rows, duration_seconds)
                        VALUES (?, ?, ?, ?, ?, ?)
                        """,
                        (
                            datetime.now().isoformat(timespec="seconds"),
                            f"leads: {os.path.basename(file_path)}",
                            len(valid_rows) + len(errors),
                            len(valid_rows),
                            len(errors),
                            duration,
                        )
                    )
                    conn.commit()
                except Exception:
                    logger.exception("Could not save lead import metadata")

            logger.info(
                "Duration=%.2fs | rows=%s | rows/s=%.1f",
                duration,
                len(valid_rows),
                len(valid_rows) / duration if duration else 0,
            )

        except Exception:
            conn.rollback()
            raise

        finally:
            conn.close()

        if dry_run:
            msg = f"OK {pluralize(len(valid_rows), 'valid lead')} analyzed."
        else:
            msg = f"OK {pluralize(len(valid_rows), 'lead')} processed: {insert_count} new."

        if duplicates_ignored > 0:
            msg += f" ({pluralize(duplicates_ignored, 'duplicate')} skipped)"
        if unlinked_leads > 0:
            msg += (
                f" ⚠ {pluralize(unlinked_leads, 'lead')} linked to unknown clients"
                " were imported without a client."
            )

        _export_failed_rows(failed_rows, "failed_leads", logger)

        if errors:
            msg += f" ⚠ {pluralize(len(errors), 'row')} rejected."
        if dry_run:
            msg += "\nDRY RUN: no data was saved to the database."

        return True, msg

    except Exception as e:
        logger.exception("Lead import failed")
        return False, f"Could not process the leads file: {e}"


def import_marketing_spend_to_db(file_path: str, dry_run: bool = False):
    import time

    start_time = time.perf_counter()
    logger = _get_importer_logger()
    logger.info("Marketing spend import started: %s", file_path)

    try:
        _check_file_size(file_path)
        ext = os.path.splitext(file_path)[1].lower()

        if ext == ".csv":
            headers, raw_rows = _read_csv_rows(file_path)
        elif ext in (".xlsx", ".xlsm"):
            headers, raw_rows = _read_excel_rows(file_path)
        else:
            return False, f"Unsupported format: {ext}. Use CSV or XLSX."

        if not headers:
            return False, "The file has no column headers."

        header_map = _build_marketing_spend_header_map(headers)

        missing = REQUIRED_MARKETING_SPEND_COLS - set(header_map.keys())
        if missing:
            missing_str = ", ".join(sorted(missing))
            found_str = ", ".join(sorted(header_map.keys()))
            return False, (
                f"Missing required columns: {missing_str}\n"
                f"Detected columns: {found_str}"
            )

        valid_rows = []
        errors = []
        failed_rows = []
        today = date.today().isoformat()

        for i, raw in enumerate(raw_rows, start=2):
            try:
                raw_source = raw.get(header_map["source"])
                source = _validate_required_string("source", raw_source)
                source = _normalize_source(source)

                raw_amount = raw.get(header_map["amount"])
                amount = _validate_positive_price("amount", raw_amount)

                raw_date = None
                if "date" in header_map:
                    raw_date = raw.get(header_map["date"])
                spend_date = _parse_date(raw_date, today)

                valid_rows.append({
                    "source": source,
                    "amount": round(amount, 2),
                    "date": spend_date,
                })

            except Exception as e:
                msg = f"Row {i}: {e}"
                errors.append(msg)
                logger.warning(msg)
                failed_rows.append({"row": i, "error": str(e), **raw})

        if not valid_rows:
            return False, "No valid rows to import."

        conn = get_connection()
        try:
            _begin_transaction(conn)
            cursor = conn.cursor()

            insert_count = 0
            for row in valid_rows:
                cursor.execute(
                    "INSERT INTO marketing_spend (source, amount, date) VALUES (?, ?, ?)",
                    (row["source"], row["amount"], row["date"]),
                )
                insert_count += 1

            if dry_run:
                conn.rollback()
                logger.info("DRY RUN: transaction rolled back, nothing saved")
            else:
                conn.commit()
                logger.info("Marketing spend transaction committed")

            duration = time.perf_counter() - start_time

            if not dry_run:
                try:
                    cursor.execute(
                        """
                        INSERT INTO import_metadata
                            (imported_at, file_name, total_rows, valid_rows, rejected_rows, duration_seconds)
                        VALUES (?, ?, ?, ?, ?, ?)
                        """,
                        (
                            datetime.now().isoformat(timespec="seconds"),
                            f"marketing_spend: {os.path.basename(file_path)}",
                            len(valid_rows) + len(errors),
                            len(valid_rows),
                            len(errors),
                            duration,
                        ),
                    )
                    conn.commit()
                except Exception:
                    logger.exception("Could not save marketing spend import metadata")

            logger.info(
                "Duration=%.2fs | rows=%s | rows/s=%.1f",
                duration,
                len(valid_rows),
                len(valid_rows) / duration if duration else 0,
            )

        except Exception:
            conn.rollback()
            raise
        finally:
            conn.close()

        if dry_run:
            msg = f"OK {pluralize(insert_count, 'spend record')} analyzed."
        else:
            msg = f"OK {pluralize(insert_count, 'spend record')} imported."

        _export_failed_rows(failed_rows, "failed_marketing_spend", logger)

        if errors:
            msg += f" ⚠ {pluralize(len(errors), 'row')} rejected."
        if dry_run:
            msg += "\nDRY RUN: no data was saved to the database."

        return True, msg

    except Exception as e:
        logger.exception("Marketing spend import failed")
        return False, f"Could not process the marketing spend file: {e}"
